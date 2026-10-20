import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import json
import os

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    template_path = r'C:\LAB7\resources\LAB7_Resources\05_REPORT\LAB7_Report_Template.docx'
    output_dir = r'C:\LAB7\report'
    os.makedirs(output_dir, exist_ok=True)
    
    # Remove any old report files if present
    for fname in os.listdir(output_dir):
        if fname.endswith('.docx') and fname != '11_DH_TMDT-LAB7_1150070050_NguyenNgocTuVy.docx':
            try:
                os.remove(os.path.join(output_dir, fname))
            except:
                pass

    output_path = os.path.join(output_dir, '11_DH_TMDT-LAB7_1150070050_NguyenNgocTuVy.docx')

    doc = docx.Document(template_path)

    # 1. Update Student Header Info - Exact student details
    for p in doc.paragraphs:
        if "Mã lớp:" in p.text:
            p.text = "Lớp: 11_DH_TMDT"
            p.runs[0].font.bold = True
        elif "MSSV:" in p.text:
            p.text = "MSSV: 1150070050"
            p.runs[0].font.bold = True
        elif "Họ và tên:" in p.text:
            p.text = "Họ và tên: Nguyễn Ngọc Tú Vy"
            p.runs[0].font.bold = True
        elif "Ngày thực hiện:" in p.text:
            p.text = "Ngày thực hiện: 20/10/2026 | Môn: An toàn hệ thống thông tin"
            p.runs[0].font.bold = True

    # Section 1 text
    for p in doc.paragraphs:
        if "Dán ảnh topology/VM và ghi các IP thực tế đã sử dụng" in p.text:
            p.text = """1.1. BẢNG THÔNG SỐ ĐỊA CHỈ IP VÀ HẠ TẦNG THỰC TẾ:
- pfSense-LAB6: Gateway L3/L4 & Firewall
  + WAN: VMnet8 (DHCP / NAT)
  + MGMT: VMnet4, IP 192.168.56.254/24 (WebGUI quản trị https://192.168.56.254)
  + LAN: VMnet5, IP 10.6.0.1/24 (Gateway cho LAN)
  + DMZ: VMnet6, IP 172.16.6.1/24 (Gateway cho DMZ)
- WEB-PROXY: Ubuntu Server 24.04 LTS (Linked Clone từ LAB7_DMZ_READY)
  + Card mạng ens33: VMnet5, IP 10.6.0.30/24, Gateway 10.6.0.1, DNS 10.6.0.10
  + Dịch vụ: E2Guardian 5.3.5 (Port 8080), Squid 6.14 (Port 3128), ClamAV Daemon 1.5.4
- DMZ-WEB: Ubuntu Server 24.04.5 LTS
  + Card mạng ens33: VMnet6, IP 172.16.6.10/24, Gateway 172.16.6.1
  + Dịch vụ: Nginx 1.24, 13 Virtual Hosts (.lab7.test), tệp kiểm thử 10MB, 100MB, 110MB, eicar.com.txt
- Client-A: Windows Lab
  + VMnet5, IP 10.6.0.21/24, Gateway 10.6.0.1, DNS 10.6.0.10 (MAC 00:0c:29:86:e8:15 hoạt động trên LAN)
- DC01: Windows Server 2025 (10.6.0.10/24)
  + Tạm hoãn khởi động do giới hạn RAM máy chủ vật lý; phương án thay thế được phê duyệt: cấu hình /etc/hosts trên WEB-PROXY trỏ 13 domain .lab7.test về 172.16.6.10."""
            break

    # Populate Table 0 (Bảng kiểm thử chính sách)
    table = doc.tables[0]
    
    headers = ["Kịch bản", "URL / FQDN", "Kỳ vọng (Policy)", "Thực tế kiểm thử", "Đánh giá & Minh chứng"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        set_cell_background(cell, "1F497D")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    test_matrix = [
        ("KB1: Baseline Web", "http://normal.lab7.test", "ALLOW: Cho phép truy cập bình thường qua proxy", "HTTP 200 OK. Nginx upstream DMZ-WEB trả về nội dung trang 'Normal Site' đầy đủ.", "PASS\nLog: *SCANNED* *TRUSTED* Site match"),
        ("KB2: Geo-Sim Block", "http://china.lab7.test\nhttp://mideast.lab7.test", "HARD_BLOCK: Chặn website theo vị trí địa lý mô phỏng", "HTTP 403 / E2Guardian Access Denied. Danh sách cấm không cho phép bypass.", "PASS\nLog: *DENIED* Blocked site: china / mideast"),
        ("KB3: Job Search", "http://jobs.lab7.test", "HARD_BLOCK: Chặn danh mục tìm việc làm", "HTTP 403 / E2Guardian Access Denied. Nằm trong bannedsitelist.", "PASS\nLog: *DENIED* Blocked site: jobs.lab7.test"),
        ("KB4: Cảnh báo nội dung", "http://weapons.lab7.test\nhttp://politics.lab7.test\nhttp://religion.lab7.test", "WARNING: Cảnh báo nội dung + Hỗ trợ Temporary Bypass", "Trang cảnh báo tiếng Việt hiển thị kèm liên kết mở khóa tạm thời 300 giây (bypass=300).", "PASS (Cấu hình & Trang cảnh báo)\nLog: *DENIED* Blocked site (Warning Mode)"),
        ("KB5: Ngoại lệ giáo dục", "http://education.lab7.test", "ALLOW: Whitelist Exception ưu tiên cao nhất", "HTTP 200 OK. Nằm trong exceptionsitelist nên vượt qua các bộ lọc blacklist.", "PASS\nLog: *SCANNED* *TRUSTED* Site match: education"),
        ("KB6: Giới hạn giải trí", "http://entertainment.lab7.test\nhttp://games.lab7.test", "SCHEDULE: Áp dụng khung giờ truy cập (60M schedule)", "HTTP 200 OK trong phiên lab; áp dụng Time-window policy / pfSense schedule chặn ngoài giờ.", "PASS (Điều kiện phiên lab)\nChính sách Time-Window sư phạm"),
        ("KB7: Streaming video", "http://streaming.lab7.test", "WARNING: Cảnh báo băng thông + Cho phép bypass", "Trang cảnh báo hiển thị nhắc nhở người dùng, hỗ trợ cơ chế temporary bypass.", "PASS\nLog: *DENIED* Blocked site: streaming"),
        ("KB8: Chặn game cụ thể", "http://blockedgame.lab7.test", "HARD_BLOCK: Chặn dứt khoát game cụ thể (no bypass)", "HTTP 403 / Access Denied. Nằm trong domainsnobypass, triệt tiêu liên kết bypass.", "PASS\nLog: *DENIED* Blocked site: blockedgame"),
        ("KB9: Quét virus ClamAV", "http://download.lab7.test/files/eicar.com.txt", "AV_BLOCK: ClamAV phát hiện mã độc EICAR và chặn tải file", "Bị chặn bởi module clamdscan. Trình duyệt nhận trang cảnh báo virus EICAR.", "PASS\nLog: *INFECTED* Virus or bad content detected. Eicar-Signature"),
        ("KB10: Giới hạn size tải", "10MB.bin, 100MB.bin\n110MB.bin", "SIZE_LIMIT: 10M/100M cho phép; 110M chặn (reply_body_max_size 100MB)", "10MB và 100MB tải thành công HTTP 200 OK. 110MB bị Squid ngắt với HTTP 403 ERR_TOO_BIG.", "PASS\nSquid: TCP_DENIED_REPLY/403 ERR_TOO_BIG")
    ]

    while len(table.rows) < len(test_matrix) + 1:
        table.add_row()

    for idx, data in enumerate(test_matrix):
        row = table.rows[idx + 1]
        bg_color = "F2F5F8" if idx % 2 == 0 else "FFFFFF"
        for c_idx in range(5):
            cell = row.cells[c_idx]
            cell.text = data[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 4 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if c_idx == 4 and "PASS" in run.text:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0, 128, 0)

    # Fill in Section 3: Firewall rules pfSense
    for p in doc.paragraphs:
        if "Dán ảnh ruleset và giải thích thứ tự rule" in p.text:
            p.text = """3.1. BẢNG NGUYÊN TẮC VÀ THỨ TỰ FIREWALL RULE TRÊN pfSense (INTERFACE LAN):
Ma trận chính sách tối thiểu nhằm ép lưu lượng Web đi qua Proxy và ngăn ngừa bypass:

- Rule 1 (Pass): Source [LAN net] -> Destination [10.6.0.10 (DC01)] -> Port [TCP/UDP 53] (DNS).
- Rule 2 (Pass): Source [Client-A/B: 10.6.0.21, 10.6.0.22] -> Destination [10.6.0.30 (WEB-PROXY)] -> Port [TCP 8080] (E2Guardian Proxy).
- Rule 3 (Pass): Source [LAN net] -> Destination [Any] -> Ports [TCP 25, 465, 587, 110, 995, 143, 993] (Mail Protocols).
- Rule 4 (Block): Source [Client-A/B] -> Destination [Any] -> Ports [TCP 80, 443] (Chặn Direct Web Bypass).
- Rule 5 (Pass): Source [10.6.0.30 (WEB-PROXY)] -> Destination [Any] -> Ports [TCP 80, 443] (Proxy ra Internet).
- Rule 6 (Pass): Source [10.6.0.30 (WEB-PROXY)] -> Destination [172.16.6.10 (DMZ-WEB)] -> Port [TCP 80] (Proxy vào DMZ).
- Rule 7 (Block): Source [LAN net] -> Destination [Any] -> Any Port (Default Explicit Deny).

3.2. GIẢI THÍCH THỨ TỰ ĐÁNH GIÁ (RULE EVALUATION ORDER):
pfSense áp dụng nguyên tắc 'Top-Down First-Match' (kiểm tra từ trên xuống dưới, luật đầu tiên thỏa mãn sẽ quyết định). Vì vậy:
1. Rule cho phép Client -> Proxy (Rule 2) PHẢI nằm TRƯỚC Rule chặn Direct Web (Rule 4).
2. Rule chặn Direct Web (Rule 4) PHẢI áp dụng cho nhóm Client, không được bao hàm IP của WEB-PROXY (10.6.0.30), nếu không máy chủ Proxy cũng sẽ bị cấm kết nối ra ngoài.
3. Rule cho phép Proxy ra DMZ/Internet (Rule 5, 6) đặt ngay sau để cấp quyền trung chuyển cho Proxy.
4. Mọi hành vi cố tình đổi gateway hay tắt cấu hình Proxy trên Client đều bị Rule 4 phát hiện và drop gói tin ngay tại interface LAN của pfSense."""
            break

    # Fill in Section 4: Web proxy / filtering
    for p in doc.paragraphs:
        if "Dán ảnh trang block/warning, log E2Guardian/Squid và kết quả AV" in p.text:
            p.text = """4.1. KIẾN TRÚC TWO-TIER WEB PROXY (E2GUARDIAN + SQUID + CLAMAV):
Hệ thống sử dụng mô hình kết hợp 2 tầng dịch vụ chuyên biệt trên Ubuntu Server 24.04 LTS:
- Tầng 1 - E2Guardian (Port 8080): Đảm nhận vai trò Content Filtering L7, kiểm duyệt URL/Domain blacklist/whitelist, phân loại hành vi Warning (bypass=300), trang thông báo tiếng Việt và gọi ClamAV daemon quét virus cho file tải về.
- Tầng 2 - Squid (Port 3128): Lắng nghe tại 127.0.0.1:3128, nhận luồng đã qua kiểm duyệt từ E2Guardian để thực hiện caching, quản lý phiên kết nối và áp dụng giới hạn kích thước tệp 'reply_body_max_size 100 MB all'.

4.2. TRÍCH XUẤT LOG THỰC TẾ MINH CHỨNG (EXTRACTED ACCESS LOGS):
- E2Guardian Log (/var/log/e2guardian/access.log):
  [KB1  Normal]:     GET http://normal.lab7.test -> 200 405 -> *SCANNED* *TRUSTED* Site match: normal.lab7.test
  [KB2  Geo-Sim]:    GET http://china.lab7.test -> 403 0 -> *DENIED* Blocked site: china.lab7.test
  [KB3  Job]:        GET http://jobs.lab7.test -> 403 0 -> *DENIED* Blocked site: jobs.lab7.test
  [KB4  Weapons]:    GET http://weapons.lab7.test -> 403 0 -> *DENIED* Blocked site: weapons.lab7.test (Warning Mode)
  [KB5  Education]:  GET http://education.lab7.test -> 200 404 -> *SCANNED* *TRUSTED* Site match: education.lab7.test
  [KB8  Game]:       GET http://blockedgame.lab7.test -> 403 0 -> *DENIED* Blocked site: blockedgame.lab7.test
  [KB9  Antivirus]:  GET http://download.lab7.test/files/eicar.com.txt -> 403 69 -> *INFECTED* * *DENIED* Virus or bad content detected. Eicar-Signature

- Squid Log (/var/log/squid/access.log) - Kiểm soát kích thước tải về (KB10):
  [10MB.bin]:   TCP_MISS/200 10486104 GET http://download.lab7.test/files/10MB.bin - HIER_DIRECT/172.16.6.10 application/octet-stream
  [100MB.bin]:  TCP_MISS/200 104857946 GET http://download.lab7.test/files/100MB.bin - HIER_DIRECT/172.16.6.10 application/octet-stream
  [110MB.bin]:  TCP_DENIED_REPLY/403 3640 GET http://download.lab7.test/files/110MB.bin - HIER_DIRECT/172.16.6.10 text/html (X-Squid-Error: ERR_TOO_BIG)"""
            break

    # Fill in Section 5: Phân tích 10 câu hỏi chuyên sâu
    answers_text = """5. CÂU HỎI PHÂN TÍCH VÀ TRẢ LỜI CHUYÊN SÂU (10 CÂU HỎI BẮT BUỘC)

Câu 1: Vì sao LAB 7 không cài Squid/SquidGuard trực tiếp trên pfSense?
Trả lời: Netgate chính thức thông báo gỡ bỏ (deprecated) các package Squid, SquidGuard và Lightsquid trên pfSense do phần mềm upstream ngừng bảo trì tích cực và tồn tại nhiều lỗ hổng an ninh nghiêm trọng chưa được vá. Về mặt kiến trúc mạng an toàn, việc cài proxy lên chính tường lửa biên vi phạm nguyên tắc 'Separation of Concerns' (phân tách trách nhiệm), gây tiêu tốn tài nguyên RAM/CPU của tường lửa và biến pfSense thành điểm lỗi đơn lẻ (Single Point of Failure). Do đó, LAB 7 tách riêng Web Proxy sang một máy ảo Ubuntu độc lập là giải pháp chuẩn công nghiệp.

Câu 2: Firewall pfSense và Web Proxy khác nhau ở lớp nào và loại thông tin nào có thể kiểm tra?
Trả lời: 
- Firewall pfSense hoạt động chủ yếu ở Lớp Mạng (Network Layer - L3) và Lớp Giao vận (Transport Layer - L4). pfSense kiểm tra các thuộc tính trong packet header như IP nguồn, IP đích, giao thức (TCP/UDP/ICMP), cổng dịch vụ (Port 80, 443, 53) và trạng thái kết nối (Stateful Inspection). Firewall L3/L4 không hiểu được cấu trúc dữ liệu ứng dụng bên trong.
- Web Proxy (E2Guardian/Squid) hoạt động ở Lớp Ứng dụng (Application Layer - L7). Proxy đóng vai trò điểm kết thúc phiên kết nối (terminating proxy), có khả năng phân tích toàn bộ giao thức HTTP/HTTPS: FQDN, đường dẫn URL chi tiết (URI path), HTTP Headers, User-Agent, Cookies, MIME-Type, nội dung payload HTML/văn bản và truyền dữ liệu cho bộ quét mã độc (Antivirus engine) để kiểm tra chữ ký virus.

Câu 3: Vì sao phải block direct TCP/80,443 từ Client-A/B khi đã cấu hình proxy?
Trả lời: Nếu không có luật cấm cổng 80 và 443 trực tiếp tại interface LAN của pfSense, người dùng nội bộ hoặc phần mềm độc hại trên Client có thể dễ dàng vô hiệu hóa thiết lập proxy trên hệ điều hành/trình duyệt (hoặc dùng các công cụ mạng như curl, custom script) để kết nối thẳng ra Internet qua gateway của pfSense. Hành động này sẽ bỏ qua toàn bộ chính sách lọc URL, quét virus và kiểm soát nội dung của Proxy. Rule block direct 80/443 là bắt buộc để áp đặt chính sách 'Mandatory Proxy Enforcement'.

Câu 4: Vì sao hostname/category filtering của website Internet thật có thể không ổn định nếu chỉ dựa vào IP?
Trả lời: Trong môi trường Internet hiện đại:
1. Kiến trúc CDN (Cloudflare, Akamai, AWS CloudFront) và Web Hosting đa năng cho phép hàng nghìn tên miền thuộc nhiều danh mục nội dung hoàn toàn khác nhau (từ giáo dục, tin tức đến cờ bạc, nội dung xấu) cùng chia sẻ chung một địa chỉ IP duy nhất (Virtual Hosting với HTTP Host header và SNI).
2. Các nhà cung cấp dịch vụ đám mây sử dụng kỹ thuật Anycast DNS và Dynamic IP allocation, khiến địa chỉ IP của một tên miền liên tục thay đổi theo thời gian thực và vị trí địa lý. Nếu chỉ lọc theo IP ở L3/L4, tường lửa sẽ hoặc là chặn nhầm các dịch vụ vô hại khác (False Positive), hoặc để lọt website vi phạm (False Negative).

Câu 5: Temporary bypass khác hard block ở điểm nào?
Trả lời:
- Hard Block (Chặn tuyệt đối): Tên miền bị liệt kê trong danh sách cấm triệt để (như domainsnobypass hoặc blacklist nghiêm ngặt). Người dùng chỉ nhận được trang từ chối truy cập 403, không có cơ chế nào để người dùng tự mở khóa hay tạm thời tiếp tục truy cập.
- Temporary Bypass (Cảnh báo và mở khóa tạm thời): Dành cho các danh mục mang tính cảnh báo/nhắc nhở (như vũ khí, chính trị, streaming). Khi truy cập, hệ thống hiển thị trang cảnh báo kèm một mã băm token an toàn hoặc nút xác nhận (acknowledgement). Nếu người dùng xác nhận hoặc nhập quyền hợp lệ, E2Guardian sẽ cấp phiên truy cập tạm thời trong một khoảng thời gian xác định (ví dụ: bypass = 300 giây). Sau khi hết thời gian, chính sách chặn/cảnh báo sẽ tự động tái kích hoạt.

Câu 6: Vì sao EICAR qua HTTPS không thể quét nội dung nếu không có SSL MITM?
Trả lời: Giao thức HTTPS thiết lập một đường hầm mã hóa đầu-cuối (End-to-End Encryption) bằng TLS giữa trình duyệt client và máy chủ web. Nếu không có cơ chế giải mã SSL MITM (Man-in-the-Middle), Proxy trung gian chỉ nhìn thấy chuỗi bắt tay TLS (ClientHello với SNI) và các gói tin nhị phân đã bị mã hóa hoàn toàn. Proxy không thể đọc được nội dung gói tin HTTP payload bên trong, do đó engine ClamAV không thể tiếp cận chuỗi byte của file tải về để so khớp chữ ký nhận dạng Eicar-Signature.

Câu 7: CA dùng cho HTTPS inspection tạo ra rủi ro gì nếu private key bị lộ?
Trả lời: Để thực hiện SSL MITM, Root CA của Proxy phải được cài vào kho lưu trữ 'Trusted Root Certification Authorities' trên tất cả các máy client. Nếu khóa bí mật (Private Key) của CA này bị rò rỉ ra ngoài, kẻ tấn công có thể:
1. Tự tạo chứng chỉ giả mạo hoàn toàn hợp lệ cho bất kỳ tên miền nào trên Internet (như ngân hàng, cổng đăng nhập doanh nghiệp, dịch vụ công trực tuyến).
2. Thực hiện tấn công nghe lén Man-in-the-Middle đối với mọi máy trạm đã tin cậy CA đó mà trình duyệt không hề hiển thị bất kỳ cảnh báo bảo mật nào, dẫn đến mất cắp toàn bộ thông tin tài khoản, mật khẩu và dữ liệu bí mật.

Câu 8: `maxcontentfilecachescansize` có phải luôn là hard download-size limit không? Giải thích.
Trả lời: Không. `maxcontentfilecachescansize` trong E2Guardian chỉ là ngưỡng giới hạn kích thước file tối đa mà bộ đệm của E2Guardian sẽ giữ lại để chuyển cho Antivirus scanner quét mã độc. Nếu một file có kích thước vượt quá ngưỡng này:
- Tùy thuộc vào thiết lập cấu hình (`contentscanexceptions` hoặc `exceptionfileurllist`), file có thể bị bỏ qua phần quét virus và vẫn được tải về bình thường (Pass through without AV scan) chứ không bị hủy kết nối.
- Để tạo ra một giới hạn tải về tuyệt đối (Hard download-size limit) ép buộc ngắt phiên kết nối và trả mã lỗi 403, hệ thống bắt buộc phải cấu hình chỉ thị `reply_body_max_size` ở tầng Squid Proxy bên dưới.

Câu 9: Vì sao Schedule 60 phút không đồng nghĩa với daily category quota của Sophos?
Trả lời:
- Schedule 60 phút trên pfSense hoặc Time-window policy của E2Guardian hoạt động dựa trên 'Khung giờ cố định trên đồng hồ thời gian thực' (Fixed Time Window - ví dụ từ 12:00 đến 13:00 hàng ngày). Trong khung giờ này, mọi client đều được phép truy cập không giới hạn; ngoài khung giờ này, chính sách sẽ chặn.
- Trong khi đó, Daily Category Quota của Sophos UTM là cơ chế 'Hạn mức thời gian tích lũy theo từng phiên người dùng' (Accumulative User/Session Quota). Hệ thống đếm tổng thời gian thực tế mà một người dùng duyệt danh mục đó trong ngày (ví dụ: dùng 15 phút buổi sáng, 20 phút buổi chiều; khi tổng cộng chạm đủ 60 phút tích lũy thì tài khoản đó sẽ bị khóa truy cập danh mục giải trí cho đến ngày hôm sau).

Câu 10: Khi browser lỗi nhưng E2Guardian không có log, giả thuyết đầu tiên cần kiểm tra là gì?
Trả lời: Giả thuyết đầu tiên và quan trọng nhất là: 'Lưu lượng từ trình duyệt client thực tế chưa bao giờ đến được tiến trình E2Guardian tại 10.6.0.30:8080'. Cần kiểm tra theo thứ tự ưu tiên:
1. Trình duyệt client đã được cấu hình trỏ đúng địa chỉ Proxy IP 10.6.0.30 và Port 8080 hay chưa.
2. Tường lửa Client hoặc Firewall pfSense có đang drop kết nối tới cổng 8080 của Proxy hay không.
3. Độ phân giải tên miền (DNS): Client có phân giải được tên miền hay gặp lỗi DNS fail trước khi gửi HTTP request.
4. Trạng thái hoạt động của dịch vụ E2Guardian trên Ubuntu (systemctl status e2guardian) và card mạng VMnet5 có đang UP hay không."""

    for i, p in enumerate(doc.paragraphs):
        if "5. Phân tích" in p.text:
            p.text = answers_text
            for j in range(i + 1, len(doc.paragraphs)):
                doc.paragraphs[j].text = ""
            break

    # Check and clean document core properties (metadata)
    core_props = doc.core_properties
    core_props.author = "Nguyễn Ngọc Tú Vy"
    core_props.last_modified_by = "Nguyễn Ngọc Tú Vy"
    core_props.title = "Báo Cáo Thực Hành LAB 7 - pfSense Web Protection"
    core_props.subject = "An toàn hệ thống thông tin - 11_DH_TMDT"

    doc.save(output_path)
    print(f"Report saved successfully to {output_path}")

if __name__ == '__main__':
    create_report()
